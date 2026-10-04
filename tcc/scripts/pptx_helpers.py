# -*- coding: utf-8 -*-
"""
pptx_helpers.py
===============
Blocos de montagem dos slides (retângulo, texto, bullets, KPI, figura, cabeçalho)
e a paleta do projeto: cinza para o contexto, azul para o dado, vermelho só no
número que se quer destacar (Knaflic, SWD-40/42).

Existe para que o deck executivo não precise duplicar os helpers do deck de
defesa — `scripts/geradores/gerar_apresentacao_pptx.py` monta os slides no
corpo do módulo e, por isso, não pode ser importado.
"""
from __future__ import annotations

from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as In
from pptx.util import Pt

# ── Paleta ────────────────────────────────────────────────────────────────────
C_DARK = RGBColor(0x1F, 0x38, 0x64)     # azul escuro: cabeçalhos
C_BLUE = RGBColor(0x15, 0x65, 0xC0)     # azul: o dado em foco
C_RED = RGBColor(0xB7, 0x1C, 0x1C)      # vermelho: só o número da frase-síntese
C_GRAY = RGBColor(0x61, 0x61, 0x61)     # cinza: contexto
C_GRAY2 = RGBColor(0x9E, 0x9E, 0x9E)
C_LGRAY = RGBColor(0xF5, 0xF5, 0xF5)    # fundo de caixa
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK = RGBColor(0x21, 0x21, 0x21)
C_AZUL_CLARO = RGBColor(0xBB, 0xDE, 0xFB)

W = In(13.33)   # 16:9
H = In(7.50)


def add_rect(slide, l, t, w, h, fill_rgb=None, line_rgb=None, line_pt=0):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.line.width = Pt(line_pt) if line_pt else 0
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb and line_pt:
        shape.line.color.rgb = line_rgb
    shape.shadow.inherit = False
    return shape


def add_text(slide, text, l, t, w, h, font_size=20, bold=False, italic=False,
             color=C_BLACK, align=PP_ALIGN.LEFT, wrap=True, font_name="Calibri"):
    caixa = slide.shapes.add_textbox(l, t, w, h)
    caixa.text_frame.word_wrap = wrap
    p = caixa.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font_name
    return caixa


def add_multiline(slide, linhas, l, t, w, h, font_size=16, color=C_BLACK,
                  bold=False, align=PP_ALIGN.LEFT, font_name="Calibri",
                  space_after=3):
    caixa = slide.shapes.add_textbox(l, t, w, h)
    caixa.text_frame.word_wrap = True
    primeiro = True
    for texto in linhas:
        p = caixa.text_frame.paragraphs[0] if primeiro else caixa.text_frame.add_paragraph()
        primeiro = False
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = texto
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name
    return caixa


def bullets(slide, itens, l, t, w, h, marcador="▸", font_size=16,
            color=C_BLACK, cor_marcador=C_BLUE):
    caixa = slide.shapes.add_textbox(l, t, w, h)
    caixa.text_frame.word_wrap = True
    primeiro = True
    for item in itens:
        p = caixa.text_frame.paragraphs[0] if primeiro else caixa.text_frame.add_paragraph()
        primeiro = False
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = marcador + "  "
        r1.font.size = Pt(font_size)
        r1.font.color.rgb = cor_marcador
        r1.font.name = "Calibri"
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(font_size)
        r2.font.color.rgb = color
        r2.font.name = "Calibri"
    return caixa


def add_img(slide, caminho, l, t, w, h=None):
    p = Path(caminho)
    if not p.exists():
        add_text(slide, f"[figura ausente: {p.name}]", l, t, w, In(0.6),
                 font_size=11, color=C_GRAY, italic=True)
        print(f"  [AVISO] figura ausente: {p.name}")
        return
    if h:
        slide.shapes.add_picture(str(p), l, t, w, h)
    else:
        slide.shapes.add_picture(str(p), l, t, w)


def header_bar(slide, titulo, subtitulo=None):
    """Título de ação (SWD-55): a primeira linha afirma o achado."""
    add_rect(slide, 0, 0, W, In(1.10), fill_rgb=C_DARK)
    add_text(slide, titulo, In(0.35), In(0.10), In(12.6), In(0.55),
             font_size=25, bold=True, color=C_WHITE)
    if subtitulo:
        add_text(slide, subtitulo, In(0.35), In(0.66), In(12.6), In(0.38),
                 font_size=14, color=C_AZUL_CLARO)


def kpi(slide, rotulo, valor, nota, l, t, w, h, cor_valor=C_BLUE):
    """Cartão de número grande: rótulo em cima, número no meio, nota embaixo."""
    add_rect(slide, l, t, w, h, fill_rgb=C_LGRAY, line_rgb=C_GRAY2, line_pt=0.75)
    add_text(slide, rotulo, l + In(0.12), t + In(0.12), w - In(0.24), In(0.5),
             font_size=13, color=C_GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, valor, l + In(0.12), t + In(0.55), w - In(0.24), In(0.9),
             font_size=40, bold=True, color=cor_valor, align=PP_ALIGN.CENTER)
    add_multiline(slide, nota if isinstance(nota, list) else [nota],
                  l + In(0.12), t + In(1.5), w - In(0.24), h - In(1.6),
                  font_size=12.5, color=C_BLACK, align=PP_ALIGN.CENTER)


def faixa_final(slide, texto, cor=C_DARK):
    """Faixa de fechamento: a mensagem do slide em uma frase."""
    add_rect(slide, In(0.3), In(6.55), In(12.73), In(0.62), fill_rgb=cor)
    add_text(slide, texto, In(0.5), In(6.62), In(12.33), In(0.5),
             font_size=15, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)


def rodape(slide, numero, fonte="PNAD Contínua 2016–2025 (IBGE) · elaboração própria"):
    add_text(slide, fonte, In(0.35), In(7.12), In(10), In(0.3),
             font_size=10, color=C_GRAY2)
    add_text(slide, str(numero), In(12.6), In(7.12), In(0.5), In(0.3),
             font_size=10, color=C_GRAY2, align=PP_ALIGN.RIGHT)
