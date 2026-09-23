# -*- coding: utf-8 -*-
"""
gerar_word_enxuto.py
====================
Converte o relatório enxuto (`relatorio_tcc_enxuto.tex`) em .docx, mantendo a
mesma fonte única de números do PDF — o Word é uma *renderização* do mesmo
LaTeX, nunca um texto paralelo (foi um texto paralelo desatualizado que deixou
o antigo relatorio_tcc.docx fora do escopo do núcleo de 4 métodos).

Etapas:
  1. expande os `\\input{}` (tabelas de outputs/tables) no próprio .tex;
  2. normaliza o que o pandoc não entende da classe abntex2cite
     (`\\citeonline`, `\\resizebox`, `\\subfigure`, caixas de destaque);
  3. chama o pandoc (pypandoc_binary) com as figuras resolvidas e as
     referências do .bib.

Saída: entregaveis/<nome de entrega>.docx e o PDF compilado ao lado
Uso:   python tcc/scripts/gerar_word_enxuto.py
"""
from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TEX = ROOT / "relatorio_tcc_enxuto.tex"
BIB = ROOT / "relatorio_tcc.bib"
FIGS = ROOT / "outputs" / "figures"
# nome de entrega (o "enxuto" era jargao interno de quando havia duas versoes)
NOME_ENTREGA = "TCC_Ricardo_Calheiros_MBA_USP_Esalq"
OUT = ROOT / "entregaveis" / (NOME_ENTREGA + ".docx")
OUT_PDF = ROOT / "entregaveis" / (NOME_ENTREGA + ".pdf")


def expandir_inputs(texto: str, base: Path, profundidade: int = 0) -> str:
    """Substitui \\input{arquivo} pelo conteúdo, recursivamente (as tabelas)."""
    if profundidade > 5:
        return texto

    def _sub(m: re.Match) -> str:
        alvo = m.group(1)
        p = base / (alvo if alvo.endswith(".tex") else alvo + ".tex")
        if not p.exists():
            print(f"  [AVISO] \\input não encontrado: {alvo}")
            return ""
        return expandir_inputs(p.read_text(encoding="utf-8"), base, profundidade + 1)

    return re.sub(r"\\input\{([^}]+)\}", _sub, texto)


def desembrulhar_resizebox(texto: str) -> str:
    """`\\resizebox{..}{..}{ TABELA }` → `TABELA`.

    O Word não tem mancha fixa; e, se o comando sobrar, o pandoc descarta a
    tabela inteira (foi assim que a tabela interseccional saiu vazia). A chave
    de fechamento é achada por contagem, não por regex, porque há chaves
    aninhadas no corpo do tabular.
    """
    padrao = re.compile(r"\\resizebox\{[^{}]*\}\{[^{}]*\}\{%?[ \t]*\n?")
    while (m := padrao.search(texto)):
        i = m.end()
        nivel = 1
        while i < len(texto) and nivel:
            if texto[i] == "{":
                nivel += 1
            elif texto[i] == "}":
                nivel -= 1
            i += 1
        # i aponta para depois da chave que fecha o \resizebox
        texto = texto[:m.start()] + texto[m.end():i - 1] + texto[i:]
    return texto


def normalizar(texto: str) -> str:
    """Tira da frente do pandoc o que é específico do abntex2cite/LaTeX visual."""
    # abntex2cite → natbib, que o pandoc entende: \citeonline{x} é citação
    # narrativa ("Pager (2007) mostra"), \cite{x} é parentética. Mapear tudo
    # para \cite quebraria as frases que usam o autor como sujeito.
    texto = re.sub(r"\\nocite\{[^}]*\}", "", texto)
    texto = re.sub(r"\\cite(online|author)\b", r"\\citet", texto)
    texto = re.sub(r"\\cite(year|yearpar)\b", r"\\citeyear", texto)
    texto = re.sub(r"\\cite(?![a-zA-Z])", r"\\citep", texto)

    # O pandoc joga o ambiente abstract em metadado e o template do docx o
    # descarta — e aqui são dois (Resumo e Abstract). Viram seções de verdade.
    nomes = iter(["Resumo", "Abstract"])
    texto = re.sub(r"\\begin\{abstract\}",
                   lambda m: "\\section*{" + next(nomes, "Resumo") + "}", texto)
    texto = texto.replace(r"\end{abstract}", "")
    texto = re.sub(r"\\renewcommand\{\\abstractname\}\{[^}]*\}", "", texto)

    # \paragraph vira cabeçalho de nível 4 numerado ("4.1.0.1") no Word; no PDF
    # é título corrido em negrito. Mantém-se o formato do PDF.
    texto = re.sub(r"\\paragraph\*?\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}", r"\\textbf{\1}", texto)

    texto = desembrulhar_resizebox(texto)

    # subfigure: fica só a imagem (o pandoc transformaria o ambiente numa tabela
    # de uma célula em volta da figura). A legenda de ação da figura-mãe permanece.
    def _sub(m: re.Match) -> str:
        img = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{[^}]+\}", m.group(1))
        return img.group(0) if img else ""

    texto = re.sub(r"\\begin\{subfigure\}(?:\[[^\]]*\])?\{[^}]*\}(.*?)\\end\{subfigure\}",
                   _sub, texto, flags=re.S)

    # o bloco do logo depende de arquivo local e vira erro silencioso
    texto = re.sub(r"\\IfFileExists\{logo_esalq\.pdf\}\{[^{}]*\}\{[^{}]*\}", "", texto)

    # regras horizontais decorativas e espaçamentos que o Word não usa
    texto = re.sub(r"\\noindent\\rule\{[^}]*\}\{[^}]*\}", "", texto)
    texto = re.sub(r"\\vspace\*?\{[^}]*\}|\\hspace\*?\{[^}]*\}", "", texto)

    # caminhos de figura absolutos em relação à raiz (o \graphicspath não viaja)
    def _fig(m: re.Match) -> str:
        opts, nome = m.group(1) or "", m.group(2)
        if "/" not in nome:
            nome = f"outputs/figures/{nome}"
        return f"\\includegraphics{opts}{{{nome}}}"

    texto = re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", _fig, texto)

    # o pandoc não resolve extensão implícita: aponta o .png de fato
    def _ext(m: re.Match) -> str:
        nome = m.group(2)
        if not Path(nome).suffix:
            for ext in (".png", ".pdf", ".jpg"):
                if (ROOT / (nome + ext)).exists():
                    nome += ext
                    break
        return f"\\includegraphics{m.group(1) or ''}{{{nome}}}"

    return re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", _ext, texto)


def main() -> int:
    try:
        import pypandoc
    except ImportError:
        print("ERRO: pypandoc não instalado. Rode: pip install pypandoc_binary")
        return 1

    if not TEX.exists():
        print(f"ERRO: {TEX} não existe — rode antes tcc/scripts/gerar_relatorio_enxuto.py")
        return 1

    print(f"Lendo {TEX.name}…")
    texto = TEX.read_text(encoding="utf-8")
    texto = expandir_inputs(texto, ROOT)
    n_tab = texto.count(r"\begin{tabular}")
    n_fig = len(re.findall(r"\\includegraphics", texto))
    texto = normalizar(texto)
    print(f"  {n_tab} tabelas e {n_fig} figuras no corpo expandido")

    with tempfile.NamedTemporaryFile("w", suffix=".tex", delete=False,
                                     encoding="utf-8", dir=str(ROOT)) as fh:
        fh.write(texto)
        tmp = Path(fh.name)

    args = [
        "--from=latex",
        "--to=docx",
        f"--resource-path={ROOT}",
        "--toc", "--toc-depth=2",
        "--number-sections",
        "--wrap=preserve",
    ]
    if BIB.exists():
        args += ["--citeproc", f"--bibliography={BIB}"]
        csl = Path(__file__).with_name("abnt.csl")     # NBR 6023/10520, como no PDF
        if csl.exists():
            args.append(f"--csl={csl}")
        else:
            print("  [AVISO] abnt.csl ausente — citações sairão no estilo padrão")

    OUT.parent.mkdir(exist_ok=True)
    try:
        print("Convertendo com pandoc…")
        pypandoc.convert_file(str(tmp), "docx", format="latex",
                              outputfile=str(OUT), extra_args=args)
    finally:
        tmp.unlink(missing_ok=True)

    from docx import Document                      # conferência do que saiu
    doc = Document(str(OUT))
    tabelas = len(doc.tables)
    imagens = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
    paragrafos = sum(1 for p in doc.paragraphs if p.text.strip())
    # o PDF é compilado na raiz sob o nome do build; a entrega leva o mesmo
    # nome do .docx, para o orientador receber um par coerente
    pdf_build = ROOT / "relatorio_tcc_enxuto.pdf"
    if pdf_build.exists():
        try:
            OUT_PDF.write_bytes(pdf_build.read_bytes())
            print(f"PDF copiado -> {OUT_PDF.relative_to(ROOT)}  "
                  f"({OUT_PDF.stat().st_size // 1024} KB)")
        except PermissionError:
            print(f"[AVISO] {OUT_PDF.name} está aberto; a cópia do PDF não foi atualizada")
    else:
        print("[AVISO] relatorio_tcc_enxuto.pdf não existe — compile o LaTeX antes")

    print(f"\nOK -> {OUT.relative_to(ROOT)}  ({OUT.stat().st_size // 1024} KB)")
    print(f"     {paragrafos} parágrafos, {tabelas} tabelas, {imagens} imagens")
    if tabelas < 8 or imagens < 6:
        print("     [AVISO] menos tabelas/figuras que o esperado — conferir o docx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
