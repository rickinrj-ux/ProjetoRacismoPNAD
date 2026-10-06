# -*- coding: utf-8 -*-
"""Checklist final: cada item cotado nos anexos do manual (p. 61 a 66)."""
import re
import zipfile
from pathlib import Path

from docx import Document

ALVO = Path(r"C:\Users\user\Documents\ProjetoRacismoPNAD\entregaveis"
            r"\TCC_Ricardo_Calheiros_MBA_USP_Esalq.docx")
d = Document(str(ALVO))
z = zipfile.ZipFile(str(ALVO))
xml = z.read("word/document.xml").decode("utf8")
cab = z.read("word/header1.xml").decode("utf8")
txt = "\n".join(p.text for p in d.paragraphs)

SEC = {"Resumo", "Abstract", "Considerações Iniciais", "Conclusão", "Referências",
       "Implementação de Algoritmo(s) de Machine Learning",
       "Resultados e Discussão"}


def fmt(nome):
    for p in d.paragraphs:
        if p.text.strip() == nome:
            return p
    return None


def prox_corpo(nome):
    achou = False
    for p in d.paragraphs:
        if achou and len(p.text.strip()) > 80:
            return p
        if p.text.strip() == nome:
            achou = True
    return None




def _quebra_entre(doc, antes: str, depois: str) -> bool:
    """Existe quebra de página entre os dois trechos?

    É o que faltava: sem quebra, o Resumo sobe para a folha de rosto, e nenhuma
    propriedade de parágrafo denuncia isso.
    """
    achou_antes = False
    for par in doc.paragraphs:
        t = par.text.strip()
        if achou_antes:
            if 'w:type="page"' in par._p.xml:
                return True
            if t.startswith(depois):
                return False
        if antes in t:
            achou_antes = True
    return False


sec = d.sections[0]
checks = []
a = checks.append

# ── folha de rosto (p. 61) ───────────────────────────────────────────────────
a(("Margens 2,5 cm nos quatro lados",
   all(abs(getattr(sec, m).cm - 2.5) < 0.05
       for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"))))
a(("Cabeçalho preenchido (sem lacunas do modelo)", "_____" not in cab))
a(("Cabeçalho em Arial 8", 'w:sz w:val="16"' in cab))
# imagem de verdade: blip com embed. Procurar só por "drawing" dava OK
# para o conector reto do template, e o logo faltava no Word.
a(("Logo do programa no cabeçalho", "r:embed" in cab and "<a:blip" in cab))
a(("Numeração desde a folha de rosto",
   'w:type="first"' not in xml or "titlePg" not in xml))
a(("Autores separados por ponto e vírgula", "Calheiros1*; Edilson" in txt))
a(("Filiação do autor com e-mail correspondente",
   "E-mail autor correspondente:" in txt))

# ── título, resumo e abstract (p. 62) ────────────────────────────────────────
for nome in ("Resumo", "Abstract"):
    p = fmt(nome)
    a((f"'{nome}': negrito, à esquerda",
       p is not None and all(r.bold for r in p.runs if r.text.strip())
       and str(p.paragraph_format.alignment).startswith("LEFT")))
    c = prox_corpo(nome)
    a((f"Texto do {nome}: simples, sem recuo, justificado",
       c is not None and c.paragraph_format.line_spacing in (1.0, None)
       and not c.paragraph_format.first_line_indent
       and str(c.paragraph_format.alignment).startswith("JUSTIFY")))

a(("Título em inglês antes do Abstract",
   "Structural racism in the Brazilian labour market" in txt))
a(("Palavras-chave e Keywords presentes",
   "Palavras-chave" in txt and "Keywords" in txt))

# ── seções do corpo (p. 63 a 66) ─────────────────────────────────────────────
for nome in ("Considerações Iniciais", "Implementação de Algoritmo(s) de Machine Learning",
             "Resultados e Discussão", "Conclusão"):
    p = fmt(nome)
    ok = (p is not None
          and all(r.bold for r in p.runs if r.text.strip())
          and not p.paragraph_format.first_line_indent
          and str(p.paragraph_format.alignment).startswith("LEFT"))
    a((f"'{nome[:34]}': negrito, à esquerda, sem recuo", ok))

# corpo com recuo e 1,5
# o Resumo e o Abstract (simples, sem recuo) ficam antes da primeira seção do corpo:
# separá-los pela posição, e não pelas primeiras palavras, que mudam quando o texto muda
_paras = list(d.paragraphs)
_ini = next((i for i, p in enumerate(_paras) if p.text.strip() == "Considerações Iniciais"), 0)
corpo = [p for p in _paras[_ini:]
         if len(p.text.strip()) > 200 and p.text.strip() not in SEC
         and not p.text.strip().startswith(("Tabela", "Figura", "Fonte", "Nota",
                                            "Palavras-chave", "Keywords"))]
a(("Corpo: 1,5 com recuo de 1,25 cm",
   bool(corpo) and corpo[0].paragraph_format.line_spacing == 1.5
   and abs(corpo[0].paragraph_format.first_line_indent.cm - 1.25) < 0.02))

# referências
refs = [p for p in d.paragraphs if re.match(r"^[A-ZÀ-Ü][A-ZÀ-Ü\s'-]{2,}(,|;)\s",
                                            p.text.strip() or "x")]
a((f"Referências ({len(refs)}): à esquerda, simples, sem recuo",
   bool(refs) and refs[0].paragraph_format.line_spacing in (1.0, None)
   and not refs[0].paragraph_format.first_line_indent
   and str(refs[0].paragraph_format.alignment).startswith("LEFT")))


# ── paginação (item 16.1 e 16.2): o que a conferência anterior não via ───────
quebras = xml.count('w:type="page"')
a((f"Quebras de página presentes ({quebras})", quebras >= 2))

# ordem dos blocos: a folha de rosto termina antes do Resumo
def _pos(trecho: str) -> int:
    return txt.find(trecho)

pos_filiacao = _pos("E-mail autor correspondente")
pos_resumo = _pos("\nResumo")
pos_abstract = _pos("\nAbstract")
pos_intro = _pos("\nConsiderações Iniciais")   # nome do template de Implementação de ML
a(("Resumo depois da folha de rosto",
   -1 < pos_filiacao < pos_resumo))
a(("Resumo começa em página nova (quebra entre ele e a capa)",
   _quebra_entre(d, "E-mail autor correspondente", "Resumo")))
a(("Abstract depois do Resumo", pos_resumo < pos_abstract))
a(("Considerações Iniciais depois do Abstract", pos_abstract < pos_intro))

# ── tabelas, figuras e notas (Tabelas 4, 7 e 9 do manual) — 05/10/2026 ───────
_notas_xml = z.read("word/footnotes.xml").decode("utf8") if "word/footnotes.xml" in z.namelist() else ""
_corpo_notas = re.findall(r'<w:footnote (?:(?!w:type="separator"|w:type="continuationSeparator")[^>])*>.*?</w:footnote>',
                          _notas_xml, re.S)
_runs_notas = [r for f in _corpo_notas for r in re.findall(r"<w:r>.*?</w:r>|<w:r .*?</w:r>", f, re.S)
               if "<w:t" in r]
a((f"Notas de rodapé em tamanho 9 ({len(_runs_notas)} trechos)",
   bool(_runs_notas) and all('<w:sz w:val="18"/>' in r for r in _runs_notas)))
_rotulos = [p.text.strip() for p in d.paragraphs
            if re.match(r"^(Tabela|Figura)\s+\d+\.", p.text.strip())
            or p.text.strip().startswith(("Fonte:", "Nota:"))]
_com_ponto = [t[:40] for t in _rotulos if t.endswith(".") and not t.endswith("...")]
a((f"Títulos, Fontes e Notas sem ponto final ({len(_com_ponto)} com ponto)", not _com_ponto))

# ── proibições ───────────────────────────────────────────────────────────────
a(("Sem Sumário (o formato não o prevê)", "Sumário" not in txt[:3000]))
a(("Conclusão sem tabela ou figura",
   "Tabela" not in txt[txt.find("\nConclusão"):] if "\nConclusão" in txt else True))

# ── limite de páginas (template de Implementação de ML: máximo de 50, com apêndices) ──
# conta no próprio Word (automação do Office): o PDF do LaTeX sai ~5 páginas menor que o
# .docx (43 × 48 em 04/10), e contar nele deixaria passar um Word acima do limite
_np, _fonte = None, "?"
try:
    import win32com.client as _w32
    _wd = _w32.DispatchEx("Word.Application")
    _wd.Visible = False
    try:
        _doc = _wd.Documents.Open(str(ALVO) if "ALVO" in globals() else str(
            Path(__file__).resolve().parents[2] / "entregaveis" / "TCC_Ricardo_Calheiros_MBA_USP_Esalq.docx"),
            False, True)
        _np, _fonte = int(_doc.ComputeStatistics(2)), "Word"
        _doc.Close(0)
    finally:
        _wd.Quit()
except Exception:
    try:   # sem Office: o PDF do LaTeX, com folga de 5 páginas
        import pypdfium2 as _pdfium
        _pdf = Path(__file__).resolve().parents[2] / "entregaveis" / "TCC_Ricardo_Calheiros_MBA_USP_Esalq.pdf"
        _np, _fonte = len(_pdfium.PdfDocument(str(_pdf))) + 5, "PDF + 5"
    except Exception:
        pass
a((f"Até 50 páginas ({_fonte}: {_np if _np is not None else '?'})",
   _np is not None and _np <= 50))

print(f"{'ITEM':62s} SITUAÇÃO")
falhas = 0
for rot, ok in checks:
    print(f"  {rot:60s} {'OK' if ok else '<<< FALHA'}")
    falhas += not ok
print(f"\n{len(checks) - falhas}/{len(checks)} conformes")
# portão de verdade: uma não conformidade derruba o passo da fila (antes saía rc 0 e a
# centralização indevida do título do Resumo passou despercebida de 03/10 a 04/10)
import sys as _sys
_sys.exit(1 if falhas else 0)
