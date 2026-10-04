# -*- coding: utf-8 -*-
"""
gerar_tcc_normas_docx.py
========================
Converte `tcc_normas.tex` em .docx usando o **template oficial do MBA USP/Esalq**
como `--reference-doc`: os estilos (Arial 11, margens de 2,5 cm, cabeçalho,
entrelinha 1,5) vêm do template, e o conteúdo vem do nosso LaTeX.

É o que torna desnecessário refazer à mão os 331 ajustes de forma que o
formatador do Sistema de TCC havia aplicado: eles passam a ser propriedade do
template, e o documento pode ser regerado quantas vezes for preciso.

Saída: entregaveis/TCC_Ricardo_Calheiros_MBA_USP_Esalq.docx
Uso:   python tcc/scripts/gerar_tcc_normas_docx.py
"""
from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[2]
TEX = ROOT / "tcc_normas.tex"
BIB = ROOT / "relatorio_tcc.bib"
CSL = Path(__file__).with_name("abnt.csl")
SAIDA = ROOT / "entregaveis" / "TCC_Ricardo_Calheiros_MBA_USP_Esalq.docx"
SAIDA_PDF = SAIDA.with_suffix(".pdf")   # o PDF da entrega sai do MESMO .tex

# o template oficial serve de referência de estilos para o pandoc
TEMPLATES = [
    Path(r"C:\Users\user\Downloads\Template TCC - Implementação de Algoritmo(s)"
         r" de Machine Learning (251, 252) (1).docx"),
    Path(r"C:\Users\user\Downloads\Template TCC_PT (251, 252).docx"),
]



FONTE_FIGURA = "Fonte: Resultados originais da pesquisa"


def sem_subfigure(texto: str) -> str:
    """Deixa só a imagem da subfigure.

    A subfigure tem legenda própria; se sobrevivesse até a numeração, levaria o
    número que pertence à figura que a contém. E o pandoc converteria o ambiente
    numa tabela de uma célula em volta da imagem.
    """
    def _sub(m):
        larg, corpo = m.group(1), m.group(2)
        img = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{[^}]+\}", corpo)
        if not img:
            return ""
        # dentro da subfigure, \textwidth é a largura dela; fora, a da página —
        # as duas imagens a 100% partiam a figura em duas páginas no Word
        return img.group(0).replace("width=\\textwidth", f"width={larg}")

    return re.sub(r"\\begin\{subfigure\}(?:\[[^\]]*\])?\{([^}]*)\}(.*?)\\end\{subfigure\}",
                  _sub, texto, flags=re.S)


def _arg_caption(corpo: str) -> tuple[int, int] | None:
    """Delimita o argumento de \\caption por contagem de chaves."""
    i = corpo.find("\\caption{")
    if i < 0:
        return None
    j, nivel = i + len("\\caption{"), 1
    while j < len(corpo) and nivel:
        if corpo[j] == "{":
            nivel += 1
        elif corpo[j] == "}":
            nivel -= 1
        j += 1
    return i + len("\\caption{"), j - 1


def legendas_numeradas(texto: str) -> str:
    """Escreve o número na legenda e acrescenta a fonte das figuras.

    No LaTeX o contador resolve a numeração; o pandoc a descarta. Escrevendo
    "Tabela 1." e "Figura 1." no corpo da legenda, o rótulo sobrevive no .docx.
    A fonte é elemento obrigatório também nas figuras (manual, item 15.1).
    """
    cont = {"table": 0, "figure": 0}
    rotulo = {"table": "Tabela", "figure": "Figura"}

    def _numerar(m):
        amb, corpo = m.group(1), m.group(2)
        cont[amb] += 1
        pos = _arg_caption(corpo)
        if pos:
            a, b = pos
            corpo = corpo[:a] + f"{rotulo[amb]} {cont[amb]}. " + corpo[a:]
        fim = ""
        if amb == "figure" and "Fonte:" not in corpo:
            # fora do ambiente: o pandoc descarta o que vem depois da legenda
            # dentro de figure, e a norma quer a fonte como linha própria
            fim = "\n\n\\noindent " + FONTE_FIGURA + "\n"
        return "\\begin{" + amb + "}" + corpo + "\\end{" + amb + "}" + fim

    return re.sub(r"\\begin\{(table|figure)\}(.*?)\\end\{\1\}", _numerar,
                  texto, flags=re.S)


def nota_depois_da_fonte(texto: str) -> str:
    """As notas de leitura viram "Nota:" — a norma reserva esse rótulo e manda
    colocá-las depois da Fonte."""
    texto = re.sub(r"\\noindent\{?Como ler (a|o) (Figura|Tabela)~\\ref\{([^}]+)\}:\}?",
                   r"Nota: como ler a \2~\\ref{\3}:", texto)
    # o rótulo costuma vir dentro de \emph{...}: trocar só "{Como ler:}" deixava
    # "\emphNota:", comando desconhecido que o pandoc descartava junto com a palavra
    texto = re.sub(r"\\(?:emph|textit|textbf)\{Como ler:\}", r"\\emph{Nota:}", texto)
    # idem com tamanho de fonte: "\footnotesize{Como ler:}" virava "\footnotesizeNota:" e
    # a nota começava por ":" no Word; o espaço separa o comando do rótulo
    texto = re.sub(r"\\(footnotesize|small|scriptsize)\{Como ler:\}", r"\\\1 Nota: ", texto)
    return texto.replace("{Como ler:}", "Nota: ").replace("Como ler:", "Nota: ")


def um_tabular_por_legenda(texto: str) -> str:
    """Tabela com dois painéis (dois tabular num só table): o pandoc repete a
    legenda em cada um. Os painéis seguintes saem do ambiente — viram tabelas sem
    legenda logo abaixo, e a nota de leitura vem com eles."""
    def _sub(m):
        corpo = m.group(1)
        partes = corpo.split("\\end{tabular}")
        if len(partes) <= 2:
            return m.group(0)
        primeiro = partes[0] + "\\end{tabular}"
        resto = "\\end{tabular}".join(partes[1:])
        return "\\begin{table}" + primeiro + "\n\\end{table}\n" + resto
    return re.sub(r"\\begin\{table\}(.*?)\\end\{table\}", _sub, texto, flags=re.S)


def resolver_referencias(texto: str) -> str:
    """Troca \\ref{rotulo} pelo texto correspondente.

    Tabelas e figuras viram "Tabela N"/"Figura N" — a numeração já foi escrita
    na legenda. Seções viram o próprio título, porque a norma não as numera e
    não há contador a que o \\ref possa apontar.
    """
    mapa: dict[str, str] = {}

    cont = {"table": 0, "figure": 0}
    rot = {"table": "Tabela", "figure": "Figura"}
    for m in re.finditer(r"\\begin\{(table|figure)\}(.*?)\\end\{\1\}", texto, re.S):
        amb = m.group(1)
        cont[amb] += 1
        for lm in re.finditer(r"\\label\{([^}]+)\}", m.group(2)):
            mapa[lm.group(1)] = f"{rot[amb]} {cont[amb]}"

    # seções e subseções: o rótulo costuma vir na linha seguinte ao título.
    # O título longo vira o nome curto (o que vem antes dos dois-pontos), para
    # a remissão não engolir a frase que a contém.
    for m in re.finditer(r"\\(?:sub)?section\*?\{([^}]*)\}\s*\n?\s*\\label\{([^}]+)\}",
                         texto):
        titulo = m.group(1).split(":")[0].split("---")[0].strip()
        mapa[m.group(2)] = "SEC:" + titulo

    faltando: set[str] = set()

    def _sub(m):
        alvo = m.group(1)
        if alvo in mapa:
            return mapa[alvo]
        faltando.add(alvo)
        return ""                      # melhor nada do que "[rotulo]" no texto

    # consome a palavra que antecede o \\ref: o texto escreve "Tabela~\\ref{x}"
    # e a troca traria "Tabela N", produzindo "Tabela Tabela N"
    texto = re.sub(r"\b(?:Tabela|Figura)s?~?\s*\\ref\{([^}]+)\}",
                   lambda m: _sub(m), texto)
    texto = re.sub(r"\\ref\{([^}]+)\}", _sub, texto)

    # remissão a seção: "Subseção~SEC:Nome" -> "ver a seção Nome". Sem isso o
    # título fica solto dentro dos parênteses e parece parte da enumeração.
    texto = re.sub(r"(?:Sub)?[Ss]e[çc][ãa]o~?\s*SEC:", "ver a seção ", texto)
    texto = re.sub(r"\bSEC:", "seção ", texto)                 # remissões soltas
    texto = re.sub(r",\s*(ver a seção)", r"; \1", texto)        # ", ver" -> "; ver"
    if faltando:
        print(f"  [AVISO] rótulos sem destino: {sorted(faltando)}")
    return texto


_MATH_SIMPLES = {r"<": "<", r">": ">", r"-": "−", r"=": "=", r"\times": "×", r"\pm": "±", r"^2": r"\textsuperscript{2}",
                 r"R^2": r"\emph{R}\textsuperscript{2}", r"\chi^2": r"χ\textsuperscript{2}",
                 r"-2\,": "−2 "}


def numeros_como_texto(texto: str) -> str:
    """Dentro das tabelas, número entre cifrões vira texto comum.

    O pandoc converte cada $-0{,}0484$ numa equação OMML: no Word a coluna fica em
    Cambria Math, não alinha com o resto e não aceita o tamanho de fonte da tabela.
    Só a matemática trivial (número, sinal, letra solta) é convertida; símbolos com
    chapéu ou índice continuam equação.
    """
    def _math(m):
        c = m.group(1).strip()
        if c in _MATH_SIMPLES:
            return _MATH_SIMPLES[c]
        num = re.fullmatch(r"([<>]?)\s*(-?)\s*(\d+(?:\{,\}\d+)?)\s*(\\%)?", c)
        if num:
            sinal = "<" if num.group(1) == "<" else ">" if num.group(1) == ">" else ""
            return (sinal + ("−" if num.group(2) else "") +
                    num.group(3).replace("{,}", ",") + (r"\%" if num.group(4) else ""))
        if re.fullmatch(r"[A-Za-z]", c):
            return r"\emph{" + c + "}"
        return m.group(0)

    def _tab(m):
        return re.sub(r"(?<!\\)\$([^$]{1,40})\$", _math, m.group(0))
    return re.sub(r"\\begin\{tabular\}.*?\\end\{tabular\}", _tab, texto, flags=re.S)


def normalizar(texto: str) -> str:
    """Prepara o LaTeX para o pandoc, como no gerador do relatório."""
    texto = re.sub(r"\\nocite\{[^}]*\}", "", texto)
    texto = re.sub(r"\\cite(online|author)\b", r"\\citet", texto)
    texto = re.sub(r"\\cite(year|yearpar)\b", r"\\citeyear", texto)
    texto = re.sub(r"\\cite(?![a-zA-Z])", r"\\citep", texto)

    # \resizebox sobrevive mal: o pandoc descarta a tabela inteira
    padrao = re.compile(r"\\resizebox\{[^{}]*\}\{[^{}]*\}\{%?[ \t]*\n?")
    while (m := padrao.search(texto)):
        i, nivel = m.end(), 1
        while i < len(texto) and nivel:
            if texto[i] == "{":
                nivel += 1
            elif texto[i] == "}":
                nivel -= 1
            i += 1
        texto = texto[:m.start()] + texto[m.end():i - 1] + texto[i:]

    # subfigure vira só a imagem (senão o pandoc cria uma tabela em volta)
    def _sub(m: re.Match) -> str:
        img = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{[^}]+\}", m.group(1))
        return img.group(0) if img else ""

    texto = re.sub(r"\\begin\{subfigure\}(?:\[[^\]]*\])?\{[^}]*\}(.*?)\\end\{subfigure\}",
                   _sub, texto, flags=re.S)

    # \paragraph vira cabeçalho numerado no Word; no PDF é título corrido
    texto = re.sub(r"\\paragraph\*?\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}",
                   r"\\textbf{\1}", texto)

    texto = re.sub(r"\\noindent\\rule\{[^}]*\}\{[^}]*\}", "", texto)
    # \cmidrule e \cline: o pandoc não os conhece e deixava "2-3(lr)4-5" no cabeçalho
    texto = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}|\\cline\{[^}]*\}", "", texto)
    # as Referências: no PDF o abnTeX imprime o título; no pandoc o --citeproc acrescenta
    # a lista no fim, sem cabeçalho — o título vai no lugar do \bibliography
    texto = re.sub(r"\\bibliography\{[^}]*\}", r"\\section*{Referências}", texto)
    texto = re.sub(r"\\bibliographystyle\{[^}]*\}", "", texto)
    texto = re.sub(r"\\vspace\*?\{[^}]*\}|\\hspace\*?\{[^}]*\}", "", texto)

    # caminhos das figuras (o \graphicspath não viaja para o pandoc)
    def _fig(m: re.Match) -> str:
        opts, nome = m.group(1) or "", m.group(2)
        if "/" not in nome:
            nome = f"outputs/figures/{nome}"
        if not Path(nome).suffix:
            for ext in (".png", ".pdf", ".jpg"):
                if (ROOT / (nome + ext)).exists():
                    nome += ext
                    break
        return f"\\includegraphics{opts}{{{nome}}}"

    return re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", _fig, texto)


def main() -> int:
    try:
        import pypandoc
    except ImportError:
        print("ERRO: pypandoc não instalado (pip install pypandoc_binary)")
        return 1
    if not TEX.exists():
        print(f"ERRO: {TEX.name} não existe — rode antes gerar_tcc_normas.py")
        return 1

    texto = TEX.read_text(encoding="utf-8")
    # a subfigure tem legenda propria e roubaria o numero da figura-mae
    texto = sem_subfigure(texto)
    texto = legendas_numeradas(texto)
    texto = nota_depois_da_fonte(texto)
    texto = resolver_referencias(texto)
    texto = um_tabular_por_legenda(texto)
    texto = numeros_como_texto(texto)
    texto = normalizar(texto)
    with tempfile.NamedTemporaryFile("w", suffix=".tex", delete=False,
                                     encoding="utf-8", dir=str(ROOT)) as fh:
        fh.write(texto)
        tmp = Path(fh.name)

    args = ["--from=latex", "--to=docx", f"--resource-path={ROOT}",
            "--wrap=preserve"]
    ref = next((t for t in TEMPLATES if t.exists()), None)
    if ref:
        args.append(f"--reference-doc={ref}")
        print(f"estilos do template: {ref.name}")
    else:
        print("  [AVISO] template oficial não encontrado; estilos padrão do pandoc")
    if BIB.exists():
        args += ["--citeproc", f"--bibliography={BIB}"]
        if CSL.exists():
            args.append(f"--csl={CSL}")

    SAIDA.parent.mkdir(exist_ok=True)
    try:
        print("Convertendo com pandoc…")
        pypandoc.convert_file(str(tmp), "docx", format="latex",
                              outputfile=str(SAIDA), extra_args=args)
    except Exception as e:
        print(f"ERRO na conversão: {e}")
        return 1
    finally:
        tmp.unlink(missing_ok=True)

    from docx import Document
    doc = Document(str(SAIDA))
    imagens = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
    print(f"\nOK -> {SAIDA.relative_to(ROOT)}  ({SAIDA.stat().st_size // 1024} KB)")
    print(f"     {len([p for p in doc.paragraphs if p.text.strip()])} parágrafos, "
          f"{len(doc.tables)} tabelas, {imagens} imagens")

    # O PDF ao lado tem de ser o mesmo documento que o .docx. Antes vinha do
    # relatório enxuto, com outro número de páginas e outro conteúdo, porque os
    # dois geradores gravavam no mesmo nome de entrega.
    pdf_build = ROOT / "tcc_normas.pdf"
    if not pdf_build.exists():
        print("[AVISO] tcc_normas.pdf não existe — compile o LaTeX "
              "(pdflatex → bibtex → pdflatex ×2) e rode de novo para o PDF da entrega")
    elif pdf_build.stat().st_mtime < TEX.stat().st_mtime:
        print("[AVISO] tcc_normas.pdf é mais antigo que tcc_normas.tex — "
              "recompile antes de entregar; o PDF não foi copiado")
    else:
        try:
            SAIDA_PDF.write_bytes(pdf_build.read_bytes())
            print(f"     PDF da mesma fonte -> {SAIDA_PDF.relative_to(ROOT)} "
                  f"({SAIDA_PDF.stat().st_size // 1024} KB)")
        except PermissionError:
            print(f"[AVISO] {SAIDA_PDF.name} está aberto; o PDF não foi atualizado")
    return 0


if __name__ == "__main__":
    sys.exit(main())
