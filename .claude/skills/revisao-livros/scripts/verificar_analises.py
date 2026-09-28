#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
verificar_analises.py — Checagens MECÂNICAS da skill /revisao-livros.

Roda em segundos, só com a biblioteca padrão, e cobre a parte automatizável
dos critérios de MHE (Angrist & Pischke), Fávero & Belfiore, Knaflic (SWD) e
apostila Alencar (ML). O julgamento qualitativo (interpretação, linguagem,
storytelling) fica com a revisão da skill — este script só levanta bandeiras.

Uso:
    python .claude/skills/revisao-livros/scripts/verificar_analises.py            # relatório completo
    python .claude/skills/revisao-livros/scripts/verificar_analises.py --quiet    # só sumário (hook)
    python .claude/skills/revisao-livros/scripts/verificar_analises.py --arquivo X # foco em X (informativo)

Saída: markdown em tcc/revisoes/verificacao_<data>.md (a menos que --no-save) e
sumário no stdout. Exit code 0 sempre (não deve bloquear o fluxo); o número de
achados ALTO/MÉDIO aparece na última linha do stdout.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]          # raiz do projeto
TEX = ROOT / "relatorio_tcc_enxuto.tex"
TABLES = ROOT / "outputs" / "tables"
NUCLEO = [                                          # tcc/run_tcc.ps1
    "run_hlm_serie_completa.py", "run_hlm_stepup.py", "run_oaxaca_blinder.py",
    "run_ob_qr_melhorias.py", "run_rif_decomp.py", "run_glmm_glassceil.py",
    "run_se_rif_interseccional.py",
]
# Scripts de rodadas anteriores que não alimentam mais nenhuma tabela do
# relatório: continuam no repositório, mas não devem gerar achado.
LEGADOS = {"run_regressao_quantilica.py", "run_hlm_serie_s20pct.py",
           "run_heckman.py", "run_sna.py", "run_po_topsis.py", "run_kmeans.py"}
ROBUSTEZ = [
    "run_ml_shap.py", "run_konfound_evalues.py", "run_interseccionalidade.py",
    "run_vif_multicolinearidade.py", "run_hlm_vs_ols_justificacao.py",
]
GERADORES = [ROOT / "scripts" / "geradores" / "gerar_apresentacao_pptx.py",
             ROOT / "tcc" / "scripts" / "gerar_relatorio_enxuto.py"]
FORA_DO_NUCLEO = [   # métodos parqueados no branch mestrado-extenso (tcc/ESCOPO_TCC.md)
    (r"\bSNA\b|betweenness|co-resid[êe]ncia|rede[s]? de capital social|análise de redes", "SNA"),
    (r"\bTOPSIS\b|programa[çc][ãa]o linear|Pesquisa Operacional|\bPO\b", "Pesquisa Operacional"),
    (r"Cluster~?\d|K-means|[Cc]lustering [Ss]ocioecon|tipologia[s]? de vulnerabilidade", "Clustering"),  # "clustering" de SE é estatística, não método
    (r"seis metodologias|três metodologias|cinco metodologias", "contagem de métodos desatualizada"),
]
# "por causa de" e "a causa de" são locuções, não afirmação causal: ficam de fora
CAUSAL = (r"\bcomprova\w*|\bprova(m|do|da)?\s+que|efeito causal|"
          r"\bdetermina(m|nte)?\b|(?<!por )(?<!a )\bcausa(m|do|ndo)?\b")

achados: list[dict] = []


def add(nivel: str, crit: str, onde: str, msg: str, fix: str = "") -> None:
    achados.append(dict(nivel=nivel, crit=crit, onde=onde, msg=msg, fix=fix))


def read(p: Path) -> str:
    try:
        return p.read_text(encoding="utf-8", errors="replace")
    except FileNotFoundError:
        return ""


def num(s: str) -> float | None:
    s = s.replace("{,}", ".").replace(",", ".").replace("−", "-").replace("$", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def close(a: float, b: float, tol: float = 0.0015) -> bool:
    return abs(a - b) <= tol


def csv_rows(name: str) -> list[dict]:
    p = TABLES / name
    if not p.exists():
        return []
    with p.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def read_tex() -> str:
    """Lê o .tex com os \input{...} de tabelas expandidos (fonte única = csv → .tex → relatório)."""
    def sub(m: re.Match) -> str:
        p = ROOT / m.group(1)
        if p.suffix == "":
            p = p.with_suffix(".tex")
        return read(p) if p.exists() else m.group(0)
    tex = re.sub(r"\\input\{([^}]+)\}", sub, read(TEX))
    return tex.replace("−", "-").replace("{,}", ",")   # normaliza sinal e vírgula decimal


def secao_de(linha_idx: int, linhas: list[str]) -> str:
    for i in range(linha_idx, -1, -1):
        m = re.match(r"\\(sub)*section\*?\{([^}]*)\}", linhas[i])
        if m:
            return m.group(2)
    return "?"


# ──────────────────────────────────────────────────────────────────────────────
# A. Consistência numérica tex ↔ csv
# ──────────────────────────────────────────────────────────────────────────────
def check_ml(tex: str) -> None:
    rows = {r["Modelo"]: r for r in csv_rows("ml_performance.csv")}
    if not rows:
        return
    for nome, chave in (("Random Forest", "Random Forest"), ("XGBoost", "XGBoost")):
        m = re.search(rf"{nome}\}}?\s*&\s*\\?t?e?x?t?b?f?\{{?([\d.,]+)\}}?\s*&\s*\\?t?e?x?t?b?f?\{{?([\d.,]+)\}}?\s*&\s*\\?t?e?x?t?b?f?\{{?([\d.,\-]+)", tex)
        if not m:
            continue
        r2, mae, rmse = (num(m.group(1)), num(m.group(2)), num(m.group(3)))
        ref = rows[chave]
        for lab, tex_v, csv_v in (("R²", r2, num(ref["R²"])), ("MAE", mae, num(ref["MAE"])),
                                  ("RMSE", rmse, num(ref["RMSE"]))):
            if tex_v is None:
                add("ALTO", "ML-04/FAV-91", "relatorio: tab:ml_perf",
                    f"{nome}: {lab} ausente na tabela do relatório (csv = {csv_v}).",
                    "Regerar a tabela a partir de outputs/tables/ml_performance.csv.")
            elif csv_v is not None and not close(tex_v, csv_v):
                add("ALTO", "ML-04/FAV-91", "relatorio: tab:ml_perf",
                    f"{nome}: {lab} no relatório = {tex_v} vs csv = {csv_v}.",
                    "Regerar a tabela a partir de ml_performance.csv (fonte única).")
    m = re.search(r"R\^2 \\approx (\d[\d{},.]*)", tex)
    if m:
        v = num(m.group(1))
        xgb = num(rows["XGBoost"]["R²"])
        if v is not None and xgb is not None and not close(v, xgb, 0.02):
            add("ALTO", "ML-04", "relatorio: texto após tab:ml_perf",
                f"Texto diz R² ≈ {v}, mas o XGBoost tem R² = {xgb}.",
                "Alinhar o texto ao csv.")
    gap = num(rows["XGBoost"].get("gap_overfit", "") or "")
    if gap is not None and gap > 0.05:
        add("MÉDIO", "ML-05", "ml_performance.csv", f"Gap treino–teste do XGBoost = {gap:.3f} (>0,05).",
            "Reduzir max_depth/aumentar regularização.")


def check_glmm(tex: str) -> None:
    rows = csv_rows("glmm_glassceil_full.csv")
    if not rows:
        return
    ors = {num(r["OR_negro"]) for r in rows}
    extras = set()
    for f in ("glmm_glassceil_glmer.csv", "grupo_rg_glmm_ocp.csv", "grupo_rg_glmm_rs_interacao.csv",
              "glmm_odds_ratios_full.csv", "evalues_glmm.csv", "grupo_rg_4grupos_desfechos.csv",
              "interseccional_coeficientes.csv"):
        for r in csv_rows(f):
            for k, v in r.items():
                if k and re.search(r"OR|odds", k, re.I):
                    x = num(v or "")
                    if x is not None:
                        extras.add(x)
    citados = re.findall(r"OR(?:\([^)]*\))?~?\$?\s*=\s*\$?\s*(\d[\d{},.]*)", tex)
    for c in citados:
        v = num(c)
        if v is None:
            continue
        if not any(close(v, o, 0.006) for o in ors | extras if o is not None):
            atuais = ", ".join(f"{r['modelo']}={num(r['OR_negro']):.3f}" for r in rows if r["desfecho"] == "ocp_qualif")
            add("ALTO", "FAV-91/SWD-06", "relatorio", f"OR = {v} citado no texto não existe em nenhum csv de GLMM/logit "
                f"(valores atuais CBO1-4: {atuais}).",
                "Atualizar o número (provavelmente resíduo de uma rodada antiga).")


def check_qr(tex: str) -> None:
    rows = csv_rows("qr_kb_test.csv")
    if not rows:
        return
    z = num(rows[0]["z_stat"])
    z_raw = num(rows[0].get("z_stat_raw", "") or "")          # versão conservadora (sem escala m/G)
    validos = [x for x in (z, z_raw) if x is not None]
    for c in re.findall(r"Z\s*=\s*\$?\s*(-?\d[\d{},.]*)", tex):
        v = num(c)
        if v is not None and validos and not any(close(v, x, 0.05) for x in validos):
            add("ALTO", "FAV-91", "relatorio", f"Teste de heterogeneidade quantílica: texto cita Z = {v}, csv = {validos}.",
                "Alinhar ao qr_kb_test.csv.")
    frac = num(rows[0].get("boot_frac", "") or "")
    if frac is not None and frac < 1 and rows[0].get("boot_blocos", "") != "UPA":
        add("MÉDIO", "MHE-75/MHE-84", "run_regressao_quantilica / qr_kb_test.csv",
            f"SE do contraste q90−q10 por bootstrap em {frac:.0%} da amostra, sem blocos por UPA.",
            "Bootstrap em blocos (UPA) ou, no mínimo, declarar no texto que o SE é conservador por vir de subamostra.")


def check_hlm(tex: str) -> None:
    # Bloco 3: se o HLM step-up (indivíduo em UPA) existir, ele é a fonte do relatório
    stepup = csv_rows("gap_decomposicao_stepup.csv")
    if stepup:
        for r in stepup:
            b = float(r["b_negro"])
            if f"{b:.4f}" not in tex and f"{b:.4f}".replace(".", ",") not in tex                     and f"{b:.4f}".replace("-", "−").replace(".", ",") not in tex:
                add("ALTO", "FAV-91", "relatorio", f"β_negro {r['Modelo']} (step-up) = {b:.4f} não aparece no relatório.",
                    "Regerar o relatório (gerar_relatorio_tcc → gerar_relatorio_enxuto).")
        fit = csv_rows("hlm_stepup_fit.csv")
        if fit and any(r["modelo"] in ("M2", "M3") and r.get("converged", "True") == "False" for r in fit):
            add("INFO", "FAV-71", "hlm_stepup_fit.csv", "BFGS reportou converged=False em M2/M3 (tolerância de gradiente).",
                "Conferir estabilidade dos coeficientes (variam <0,1% entre degraus) ou aumentar maxiter.")
        return
    rows = csv_rows("hlm_serie_completo.csv")
    if not rows:
        return
    negro = next((r for r in rows if r.get("") == "negro"), None)
    if negro:
        for col, rotulo in (("M1_Individual", "M1"), ("M2_Localidade", "M2"), ("M4_Ocupacao", "M4")):
            m = re.match(r"(-?\d+\.\d+)", negro[col])
            if not m:
                continue
            b = float(m.group(1))
            if f"{b:.4f}" not in tex and f"{b:.4f}".replace(".", ",") not in tex:
                add("ALTO", "FAV-91", "relatorio", f"β_negro {rotulo} = {b} do csv não aparece no relatório.",
                    "Regerar tabelas/texto a partir de hlm_serie_completo.csv.")
        se_hlm = re.search(r"\((\d+\.\d+)\)", negro["M1_Individual"])
        se_ols = re.search(r"\((\d+\.\d+)\)", negro["M1_Individual_OLS"])
        linha_negro = next((l for l in tex.splitlines() if "textbf{Raça (negro)}" in l and "&" in l), "")
        publica_cluster = "[" in linha_negro          # tabela do relatório já traz SE agrupado entre colchetes
        if se_hlm and se_ols and float(se_ols.group(1)) > 3 * float(se_hlm.group(1)) and not publica_cluster:
            add("ALTO", "MHE-81/MHE-04", "hlm_serie_completo.csv",
                f"SE de β_negro: HLM (RE de UF) = {se_hlm.group(1)} vs OLS cluster-UF = {se_ols.group(1)} "
                f"({float(se_ols.group(1))/float(se_hlm.group(1)):.0f}×). O relatório publica só o menor (Moulton).",
                "Reportar SE clusterizado (UPA, e UF com correção de poucos clusters) ao lado do SE do HLM; "
                "nunca (0.0000).")
    # colinearidade UF z-scores + C(UF_str) nos OLS
    for r in rows:
        for col in ("M3_Completo_OLS", "M4_Ocupacao_OLS"):
            v = r.get(col, "")
            if re.search(r"\d{7,}", v) or "nan" in v:
                add("MÉDIO", "FAV-33/FAV-44", "hlm_serie_completo.csv",
                    f"{col}: coeficiente degenerado em '{r.get('')}' ({v[:40]}…) — z-scores de UF são colineares com C(UF_str).",
                    "Remover _UF das fórmulas OLS com dummies de UF (ou remover as dummies).")
                break
        else:
            continue
        break
    if re.search(r"AIC\s*&\s*N/D", tex):
        add("MÉDIO", "FAV-71/FAV-72", "relatorio: tab:hlm_resultados",
            "AIC/BIC 'N/D' em todos os modelos: sem LR test/critérios de informação a estratégia step-up (Fávero) fica incompleta.",
            "Reportar −2LL, AIC, BIC e LR test M0→M1→M2→M3 (REML comparável ou refit por ML).")
    if re.search(r"\(0\.0000\)", tex):
        add("BAIXO", "MHE-85/FAV-91", "relatorio: tab:hlm_resultados",
            "Erros-padrão exibidos como (0.0000).", "Mais casas decimais ou notação científica.")
    if re.search(r"interceptos fixos", tex):
        script = read(ROOT / "scripts" / "analise" / "run_hlm_serie_completa.py")
        if "C(UPA" not in script:
            add("ALTO", "FAV-70/FAV-71/MHE-56", "relatorio: subsec:hlm vs run_hlm_serie_completa.py",
                "Texto diz que a UPA entra como 'interceptos fixos (41.517 grupos)', mas o script estima "
                "intercepto aleatório de UF + covariáveis contextuais de UPA (sem C(UPA)). Com dummies de UPA, "
                "%negro-UPA seria não identificável.",
                "Reescrever a metodologia: modelo de 2 níveis (indivíduo/UF) com contexto de UPA como covariáveis; "
                "ou estimar de fato o 3º nível (run_hlm_m4 usa groups=UPA) e reportar τ²_UPA e ICC_UPA.")


def check_ob(tex: str) -> None:
    rows = {r["Componente"]: r for r in csv_rows("oaxaca_resultados.csv")}
    if not rows:
        return
    dot = num(rows["Efeito Dotacoes"]["Pct_do_gap"])
    ret = num(rows["Efeito Retornos"]["Pct_do_gap"])
    for v, lab in ((dot, "dotações"), (ret, "retornos")):
        if v is None:
            continue
        s1, s2 = f"{v:.1f}", f"{v:.1f}".replace(".", ",")
        if s1 not in tex and s2 not in tex:
            add("ALTO", "FAV-91", "relatorio", f"% {lab} do OB no csv ({v:.1f}) não aparece no texto.",
                "Regerar ob_acesso.tex/texto.")
    script = read(ROOT / "scripts" / "analise" / "run_oaxaca_blinder.py")
    if "ocp_" in script and "bootstrap" not in script.lower():
        add("MÉDIO", "MHE-81/FAV-91", "run_oaxaca_blinder.py",
            "Decomposição sem erro-padrão (nem bootstrap) no script; a legenda da tabela promete 'bootstrap 200 rep. por UF'.",
            "Confirmar de onde vêm os SE da tabela (run_ob_qr_melhorias?) e citar o script certo; bootstrap em blocos.")
    if "ocp_" in script:
        add("INFO", "MHE-25", "run_oaxaca_blinder.py",
            "OB inclui ocupação/formalidade/horas como dotações (bad controls). Correto SÓ se o texto apresentar também "
            "a versão sem ocupação e tratar 16,2% como limite inferior.", "Ver critério MHE-25 na revisão qualitativa.")


def check_gap_interno(tex: str) -> None:
    liq = {num(x) for x in re.findall(r"gap l[íi]quido[^.\n]{0,40}?(-?\d[\d{},.]*)\\%", tex)}
    res = {num(x) for x in re.findall(r"gap residual[^.\n]{0,40}?(-?\d[\d{},.]*)\\%", tex)}
    liq.discard(None); res.discard(None)
    if len(liq) > 1:
        add("ALTO", "SWD-79/FAV-91", "relatorio", f"'gap líquido' citado com valores diferentes: {sorted(liq)}.",
            "Um único número-síntese (M3 = −9,6% sem ocupação; M4 = −6,2% com ocupação) e nomear cada um.")
    if len(res) > 1:
        add("ALTO", "SWD-79/FAV-91", "relatorio", f"'gap residual' citado com valores diferentes: {sorted(res)}.",
            "Idem.")
    if liq and res and not (liq & res) and abs(min(liq) - min(res)) > 1:
        add("MÉDIO", "SWD-79", "relatorio",
            f"'gap líquido' ({sorted(liq)}) e 'gap residual' ({sorted(res)}) são usados como sinônimos com números diferentes.",
            "Fixar terminologia: líquido = M3 (sem bad controls); residual pós-ocupação = M4.")
    # Fonte única: nenhum marcador de geração pode sobreviver no .tex entregue.
    # (Em 2026-09-22 uma nota de tabela em raw string deixou '{fmt(abs(g1),1)}' chegar ao PDF.)
    vazados = set(re.findall(r"\{fmt\(.{0,40}?\)\}|@@[A-Z0-9_]+@@|\{[a-z_]+\[[^\]]{0,30}\]\}", tex))
    if vazados:
        add("ALTO", "FONTE-ÚNICA/FAV-91", "relatorio",
            f"Marcador de geração não preenchido no .tex: {sorted(vazados)[:5]}.",
            "O texto veio de uma raw string ou de um placeholder sem filler; preencher na "
            "geração e reconferir o PDF.")

    n_obs = set(re.findall(r"N\s*=\s*\$?([\d.]{7,})", tex))
    # N diferentes são esperados (cada modelo tem seus filtros); o problema é o N
    # que aparece no texto sem estar declarado em nenhuma legenda de tabela.
    legendas = " ".join(re.findall(r"\\caption\{(.*?)\}\s*\n", tex, re.S))
    n_obs = {n for n in n_obs if n not in legendas}
    if len(n_obs) > 1:
        add("BAIXO", "FAV-91", "relatorio", f"Vários N citados como 'população': {sorted(n_obs)}.",
            "Declarar o N de cada modelo (filtros diferentes) em cada tabela.")


# ──────────────────────────────────────────────────────────────────────────────
# B. Escopo e linguagem
# ──────────────────────────────────────────────────────────────────────────────
def check_escopo(linhas: list[str]) -> None:
    for pat, nome in FORA_DO_NUCLEO:
        # menções em Limitações/agenda ("versão estendida", "trabalhos futuros") são legítimas
        hits = [i + 1 for i, l in enumerate(linhas)
                if not l.lstrip().startswith("%") and re.search(pat, l)
                and "Limita" not in secao_de(i, linhas)]
        if hits:
            add("ALTO", "ESCOPO/SWD-51", "relatorio",
                f"{nome} (fora do núcleo de 4) ainda citado em {len(hits)} linha(s): {hits[:12]}{'…' if len(hits) > 12 else ''}.",
                "Remover ou mover para 'trabalhos futuros'; ajustar Hipóteses, mapa das camadas, Conclusão e título.")


def check_causal(linhas: list[str]) -> None:
    hits = []
    for i, l in enumerate(linhas):
        if l.lstrip().startswith("%"):
            continue
        sec = secao_de(i, linhas)
        if "Limita" in sec:
            continue
        if re.search(CAUSAL, l, re.I):
            # negações explícitas são o uso correto ("não constitui prova", "não é efeito
            # causal"); a negação pode estar na linha anterior, então olha-se o par de linhas
            ctx = (linhas[i - 1] + " " + l) if i else l
            if re.search(r"\bn[ãa]o\b[^.]{0,120}(" + CAUSAL + ")", ctx, re.I):
                continue
            hits.append(i + 1)
    if hits:
        add("MÉDIO", "MHE-01/MHE-23/FAV-07", "relatorio",
            f"Linguagem causal ('comprova', 'determina', 'causa', 'efeito causal') em {len(hits)} linha(s): {hits[:15]}.",
            "Trocar por 'associa-se', 'é consistente com', 'penalidade condicional'; reservar causal para a seção de limitações.")


# ──────────────────────────────────────────────────────────────────────────────
# C. Especificação nos scripts do núcleo
# ──────────────────────────────────────────────────────────────────────────────
def check_scripts() -> None:
    sem_peso: list[str] = []
    for nome in NUCLEO + ROBUSTEZ:
        p = ROOT / "scripts" / "analise" / nome
        src = read(p)
        if not src:
            add("MÉDIO", "ESCOPO", nome, "Script do núcleo/robustez não encontrado.", "")
            continue
        code = "\n".join(re.sub(r"#.*$", "", l) for l in src.splitlines())   # sem comentários
        # amostragem (regra do projeto: nunca amostrar sem confirmação)
        m = re.search(r"^SAMPLE_FRAC\s*=\s*(0\.\d+)", code, re.M)   # só atribuição real (não docstring)
        if m:
            add("ALTO", "REGRA-AMOSTRAGEM/FAV-09", nome, f"SAMPLE_FRAC = {m.group(1)} (não é população completa).",
                "Voltar a None; amostra só com confirmação explícita do usuário.")
        if re.search(r"\.sample\(", code) and "sample_frac" not in code.lower() \
                and nome not in ("run_ml_shap.py", "run_vif_multicolinearidade.py"):
            add("MÉDIO", "REGRA-AMOSTRAGEM", nome, "Uso de .sample( no script.", "Confirmar se é só para visualização.")
        # pesos amostrais
        if not re.search(r"V1028|weights=|pweight|freq_weights|var_weights", code):
            sem_peso.append(nome)
        # OLS sem cov_type. Duas isenções legítimas: o script de justificação
        # compara SE ingênuo vs. cluster de propósito, e nas decomposições o fit
        # só fornece coeficientes — o erro-padrão vem do bootstrap em blocos.
        _isento = (nome == "run_hlm_vs_ols_justificacao.py"
                   or bool(re.search(r"boot\w*.{0,60}UPA|UPA.{0,60}boot\w*|blocos? (de|por) UPA",
                                     code, re.I | re.S)))
        for mm in (re.finditer(r"smf\.ols\([^\n]*\)\s*\.fit\(\s*\)", code)
                   if not _isento else ()):
            add("MÉDIO", "MHE-22/MHE-81", nome, f"OLS com .fit() sem cov_type: `{mm.group(0)[:70]}…`",
                "cov_type='cluster' (UPA) ou HC3; regra do máximo (MHE cap. 8).")
        if "quantreg" in code and "bootstrap" not in code.lower() and "boot" not in code.lower():
            add("INFO", "MHE-75", nome, "QR sem bootstrap: SE dependem da densidade do resíduo em zero.",
                "Bootstrap em blocos por UPA (ou declarar o kernel usado).")
        if "smf.logit" in code and not re.search(r"glmer|mixedlm|BinomialBayesMixedGLM|MixedLM", code):
            # bloco 4: o GLMM real vem de scripts/R/glmm_glassceil.R (glmer); o logit-FE é robustez
            if "glmm" in nome and not (TABLES / "glmm_glassceil_glmer.csv").exists():
                add("ALTO", "FAV-77/FAV-70", nome,
                    "Script rotulado GLMM estima logit com UF como efeito fixo + HC1 (sem efeito aleatório). "
                    "O relatório descreve 'GLMM com efeito aleatório de UPA'.",
                    "Ou alimentar a tabela com scripts/R/logit_multinivel_glmm.R (glmer, ICC UPA, LR test vs logit), "
                    "ou renomear para 'logit com efeitos fixos de UF' e citar o glmer como robustez.")
        # a CV e a busca de hiperparâmetros vivem em run_ml_cv.py (bloco 8); o script de
        # produção usa a configuração escolhida lá
        if "train_test_split" in code and not re.search(r"cross_val|KFold|GridSearch|RandomizedSearch|optuna", code)                 and not (TABLES / "ml_cv_resumo.csv").exists():
            add("MÉDIO", "ML-02/ML-06", nome, "Hold-out único, sem validação cruzada nem busca de hiperparâmetros.",
                "k-fold (k=5) em subamostra para escolher max_depth/lr/n_estimators; reportar CV e hold-out.")
        if re.search(r"\.pie\(|projection=['\"]3d['\"]|twinx\(\)", code):
            add("MÉDIO", "SWD-19", nome, "Gráfico de pizza / 3D / eixo secundário detectado.", "Barras ou slopegraph.")
        if nome == "run_glmm_glassceil.py" and not re.search(r"roc_auc|AUC|hosmer|confusion", code, re.I)                 and not (TABLES / "glmm_glassceil_glmer.csv").exists():
            add("MÉDIO", "FAV-43/ML-03", nome, "Logit sem AUC/ROC, matriz de confusão (cutoff) nem Hosmer-Lemeshow.",
                "Reportar AUC + sensibilidade/especificidade no cutoff; comparar M1→M3 (roccomp).")
    if sem_peso:
        add("INFO", "MHE-30/FAV-09/FAV-36", ", ".join(sem_peso),
            "Sem peso amostral da PNAD (V1028) em nenhum modelo.",
            "Declarar no texto que as estimativas são não ponderadas (regressão amostral, não populacional) "
            "ou reestimar um modelo-chave com pesos como robustez.")
    # geradores/relatório citam método parqueado?
    for g in GERADORES:
        if g.name == "gerar_relatorio_enxuto.py":   # cita os métodos só para removê-los
            continue
        src = read(g)
        for pat, nome in FORA_DO_NUCLEO[:3]:
            if re.search(pat, src):
                add("MÉDIO", "ESCOPO", g.relative_to(ROOT).as_posix(), f"Gerador ainda referencia {nome}.",
                    "Limpar para que a regeneração não reintroduza o método.")


# ──────────────────────────────────────────────────────────────────────────────
# D. Storytelling: figuras e slides
# ──────────────────────────────────────────────────────────────────────────────
def check_visual(tex: str) -> None:
    figs = re.findall(r"\\includegraphics\[[^\]]*\]\{([^}]+)\}", tex)
    figs = [f for f in figs if "logo" not in f]
    metodos = {"HLM": r"hlm|mediacao|gap_", "Oaxaca-Blinder": r"oaxaca|ob_|cascata|waterfall_ob",
               "Regressão quantílica/RIF": r"quantreg|qr_|rif", "GLMM/logit": r"glmm|logit|odds|or_"}
    for met, pat in metodos.items():
        if not any(re.search(pat, f, re.I) for f in figs):
            add("MÉDIO", "SWD-10/SWD-18/SWD-14", "relatorio: figuras",
                f"Nenhuma figura para {met} (só tabela). Figuras atuais: {len(figs)}.",
                "Uma figura por método do núcleo: cascata (OB), coeficiente×quantil com IC (QR/RIF), "
                "barras horizontais de OR com IC (GLMM), barras do β_negro M1→M4 (HLM).")
    # SWD-55 vale para FIGURA: legenda de tabela é descritiva por convenção
    # acadêmica (a ABNT pede que identifique o conteúdo, não que argumente).
    caps = []
    for amb in re.finditer(r"\\begin\{figure\}(.*?)\\end\{figure\}", tex, re.S):
        caps += re.findall(r"\\caption\{([^}]{0,120})", amb.group(1))
    descritivos = [c for c in caps if re.match(r"\s*(Análise|Decomposição|Modelos|Desempenho|Dependence|Razões|Razoes)", c)]
    if descritivos:
        add("BAIXO", "SWD-55", "relatorio: legendas",
            f"{len(descritivos)} legenda(s) começam por rótulo descritivo em vez de afirmar o achado (ex.: '{descritivos[0][:60]}…').",
            "Título de ação na primeira frase da legenda.")


def check_slides() -> None:
    src = read(GERADORES[0])
    if not src:
        return
    # o f no prefixo é opcional: títulos com número lido de csv são f-strings
    titulos = re.findall(r'header_bar\(s,\s*f?"([^"]+)"', src)
    nums = [int(m.group(1)) for t in titulos if (m := re.match(r"(\d+)\.", t))]
    if nums and nums != list(range(nums[0], nums[0] + len(nums))):
        add("MÉDIO", "SWD-75/SWD-51", "gerar_apresentacao_pptx.py",
            f"Numeração dos slides com lacunas: {nums} (slides removidos deixaram rastro).",
            "Renumerar ou remover números; a sequência de títulos deve contar a história sozinha.")
    descritivos = [t for t in titulos if not re.search(r"[a-záéíóúç]+(a|e|em|am|ou|ia|ram)\b.*[a-z]", t.split("—")[-1].lower())
                   or re.match(r"^\d+\.\s*[\w\-/ ]+(—|-)\s*[\w ]+$", t)]
    if len(titulos) and len(descritivos) >= len(titulos) // 2:
        add("MÉDIO", "SWD-55/SWD-75", "gerar_apresentacao_pptx.py",
            f"{len(descritivos)}/{len(titulos)} títulos de slide são tópicos ('N. Método — Tema'), não frases de ação.",
            "Ex.: '5. HLM — Decomposição do Gap' → 'Morar em bairro segregado explica metade do gap racial'.")
    # SWD-42 fala de cores que *codificam dado*; branco, preto, cinzas e o fundo do tema
    # são estruturais e não contam para a saturação da paleta.
    NEUTRAS = re.compile(r"WHITE|BLACK|GRAY|GREY|DARK|LGRAY|BG|FUNDO", re.I)
    nomeadas = re.findall(r"^(C_\w+)\s*=\s*RGBColor\(", src, re.M)
    cores = len({n for n in nomeadas if not NEUTRAS.search(n)})
    if cores > 5:
        add("BAIXO", "SWD-42/SWD-40", "gerar_apresentacao_pptx.py",
            f"Paleta com {cores} cores de dado (fora as neutras).",
            "Cinza + uma cor de destaque (azul); vermelho só para o dado que se quer destacar.")



# ──────────────────────────────────────────────────────────────────────────────
# D. Entregáveis binários (.pptx/.docx) contra os csv
# ──────────────────────────────────────────────────────────────────────────────
def _col_csv(arquivo: str, coluna: str, escala: float = 1.0) -> set[float]:
    vals = set()
    for r in csv_rows(arquivo):
        try:
            vals.add(float(r[coluna]) * escala)
        except (TypeError, ValueError, KeyError):
            pass
    return vals


def _formatos(vals, casas) -> set[str]:
    saida = set()
    for v in vals:
        for d in casas:
            saida.add(f"{v:.{d}f}".replace(".", ","))
            saida.add(f"{abs(v):.{d}f}".replace(".", ","))
    return saida


def _texto_entregavel(p: Path) -> str:
    """Extrai o texto de um .pptx ou .docx; devolve '' se a lib não existir."""
    try:
        if p.suffix == ".pptx":
            from pptx import Presentation
            pr = Presentation(str(p))
            return "\n".join(sh.text_frame.text for sl in pr.slides for sh in sl.shapes
                             if sh.has_text_frame)
        from docx import Document
        d = Document(str(p))
        t = "\n".join(par.text for par in d.paragraphs)
        for tb in d.tables:
            for r in tb.rows:
                t += "\n" + " ".join(c.text for c in r.cells)
        return t
    except Exception:
        return ""


def check_entregaveis() -> None:
    """Todo número ancorado num rótulo (OR, ICC, AME, R², E-value) dentro de um
    entregável tem de existir no csv correspondente. É a checagem que pega um
    .pptx/.docx que deixou de acompanhar a reexecução das análises."""
    ent = ROOT / "entregaveis"
    if not ent.is_dir():
        return

    regras = {
        "OR": (r"OR\s*[=:≈]?\s*(\d,\d{2,3})",
               _formatos(_col_csv("glmm_glassceil_glmer.csv", "OR_negro")
                         | _col_csv("glmm_glassceil_glmer.csv", "OR_inter_superior")
                         | _col_csv("glmm_glassceil.csv", "OR_negro")
                         | _col_csv("grupo_rg_4grupos_desfechos.csv", "OR_mulher_negra")
                         | _col_csv("grupo_rg_4grupos_desfechos.csv", "OR_homem_negro")
                         | _col_csv("grupo_rg_4grupos_desfechos.csv", "OR_mulher_branca"),
                         (2, 3))),
        "ICC": (r"ICC[^=\n]{0,20}[=:]\s*(\d{1,2},\d{1,2})\s*%",
                _formatos(_col_csv("hlm_stepup_fit.csv", "icc_upa", 100)
                          | _col_csv("glmm_glassceil_glmer.csv", "ICC_UPA", 100), (0, 1, 2))),
        "AME": (r"AME\s*[=:]?\s*[−-]?\s*(\d{1,2},\d{1,2})",
                _formatos(_col_csv("glmm_glassceil_glmer.csv", "AME_pp"), (1, 2))),
        "R²": (r"R²\s*(?:de teste\s*)?[=:]?\s*(\d,\d{2,4})",
               _formatos(_col_csv("ml_performance.csv", "R²")
                         | _col_csv("ml_performance.csv", "R2_treino")
                         | _col_csv("ml_cv_resumo.csv", "cv_r2_media")
                         | _col_csv("ml_cv_resumo.csv", "teste_r2")
                         | _col_csv("ml_cv_hiperparametros.csv", "r2")
                         | _col_csv("ml_cv_hiperparametros.csv", "treino_r2"), (2, 3, 4))),
        "E-value": (r"E-value[^=\n]{0,12}[=≥]\s*(\d,\d{1,2})",
                    _formatos(_col_csv("evalues_glmm.csv", "E-value (OR)"), (1, 2))),
    }

    for p in sorted(ent.glob("*.pptx")) + sorted(ent.glob("*.docx")):
        if p.name.startswith("~$"):        # arquivo de bloqueio do Word
            continue
        txt = _texto_entregavel(p)
        if not txt:
            continue
        ruins = []
        for rotulo, (padrao, validos) in regras.items():
            if not validos:
                continue
            for m in re.finditer(padrao, txt):
                if m.group(1) not in validos:
                    ruins.append(f"{rotulo} = {m.group(1)}")
        if ruins:
            add("ALTO", "FONTE-ÚNICA/FAV-91", f"entregaveis/{p.name}",
                f"Valores sem correspondência no csv: {', '.join(sorted(set(ruins))[:6])}.",
                "O entregável não acompanhou a reexecução: regerá-lo lendo os csv "
                "(params_nucleo.py) em vez de números escritos à mão.")


# ──────────────────────────────────────────────────────────────────────────────
# E. Remissões e leitura de coeficientes em log
# ──────────────────────────────────────────────────────────────────────────────

# ── N6 — carga cognitiva (SWD-30, SWD-40, SWD-71) ──────────────────────────
NUM_CARGA = re.compile(r"(?<![\w,.])\d+(?:[.,]\d+)?")
# convenção e endereço não são carga: g.l., p-valores, IC nominal, leis, artigos
SEM_CARGA = re.compile(
    r"p\s*[<=]\s*0\{?,?\}?\d+|~?95\\%|IC~?95|\d+~?g\.l\.|\(\d+\)\s*,?\s*\$p"
    r"|Lei~?[\d.]+/\d{4}|art\.~?\d+|n[íi]vel~?\d|cap\.~?\d+|\bCBO~?1--4\b"
    r"|top~?\d+\\%|\\ref\{[^}]*\}|\\cite\w*\{[^}]*\}|20\d{2}--20\d{2}")


def _sem_notas(t: str) -> str:
    """Remove \\footnote{...} contando chaves: o conteúdo aninha vários níveis."""
    fora, i = [], 0
    while True:
        j = t.find(r"\footnote{", i)
        if j < 0:
            fora.append(t[i:])
            return "".join(fora)
        fora.append(t[i:j])
        k, prof = j + len(r"\footnote{"), 1
        while k < len(t) and prof:
            prof += (t[k] == "{") - (t[k] == "}")
            k += 1
        i = k


def _prosa(t: str) -> str:
    t = SEM_CARGA.sub(" ", _sem_notas(t))
    t = re.sub(r"\\[a-zA-Z]+\*?", " ", t)
    return re.sub(r"[{}$\\~]", " ", t)


def check_carga_cognitiva(tex: str) -> None:
    corpo = re.sub(r"\\begin\{(table|figure|tabular|equation\*?)\}.*?\\end\{\1\}",
                   " ", tex, flags=re.S)
    # no enxuto são duas seções (Resultados; Discussão e Prescrição); no normativo,
    # uma só (Resultados e Discussão). Mede-se de uma delas até a Conclusão.
    m = re.search(r"\\section\*?\{Resultados\b.*?\}(.*?)"
                  r"\\section\*?\{Conclusão\}", corpo, re.S)
    if not m:
        return
    sec = m.group(1)

    pesados, sem_pergunta = [], []
    blocos = re.split(r"\\subsection\*?\{([^}]*)\}", sec)
    for i in range(1, len(blocos), 2):
        nome, txt = blocos[i], blocos[i + 1]

        # abre com número? (SWD-71: falta o parágrafo de pergunta)
        abertura = _prosa(txt.split("\\paragraph")[0])
        # NUM_CARGA (e não `\d`) para não confundir rótulo de modelo,
        # como "M4" ou "H2", com resultado numérico
        if NUM_CARGA.search(" ".join(abertura.split()[:15])):
            sem_pergunta.append(nome[:44])

        partes = re.split(r"\\paragraph\{([^}]*)\}", txt)
        itens = [("(abertura)", partes[0])]
        itens += [(partes[k], partes[k + 1]) for k in range(1, len(partes), 2)]
        for rot, p in itens:
            c = _prosa(p)
            w, n = len(c.split()), len(NUM_CARGA.findall(c))
            if w < 40:
                continue
            d = n / w * 100
            # denso de verdade, ou muitos números mesmo num parágrafo longo
            if d > 6.0 or (n > 6 and d > 4.5):
                pesados.append(f"{nome[:22]} | {rot[:34]} ({n} em {w} = {d:.1f}%)")

    if pesados:
        add("MÉDIO", "SWD-30/SWD-40", "relatorio",
            f"{len(pesados)} parágrafo(s) de prosa com carga numérica alta "
            f"(acima de 6%, ou mais de 6 números acima de 4,5%): {pesados[:6]}.",
            "Manter na linha de leitura só o que a banca vai citar; erro-padrão, IC, "
            "estatísticas de teste e componentes de variância vivem na tabela ou em nota.")
    if sem_pergunta:
        add("BAIXO", "SWD-71", "relatorio",
            f"{len(sem_pergunta)} subseção(ões) de resultado abrem com número nas "
            f"primeiras 15 palavras: {sem_pergunta}.",
            "Abrir com uma ou duas frases que digam o que está para ser descoberto e "
            "por que importa; o número entra depois da pergunta.")

def check_remissoes() -> None:
    """Remissão que chega quebrada ao leitor.

    Nasceu de um defeito real: ao resolver \\ref em texto (a norma não numera
    seções), o título ia nu para dentro dos parênteses e o leitor o lia como
    parte da enumeração — "(OLS com efeitos fixos de UF, Inferência:
    erros-padrão agrupados, poucos clusters e pesos)".
    """
    # só os entregáveis: no .tex o \ref ainda não foi resolvido, e é o correto
    alvos: list[Path] = []
    ent = ROOT / "entregaveis"
    if ent.is_dir():
        alvos += [p for p in ent.glob("*.docx") if not p.name.startswith("~$")]

    for alvo in alvos:
        if not alvo.exists():
            continue
        txt = read(alvo) if alvo.suffix == ".tex" else _texto_entregavel(alvo)
        if not txt:
            continue
        onde = alvo.name

        crus = set(re.findall(r"\[(?:fig|tab|sec|subsec|eq):[^\]]+\]", txt))
        if crus:
            add("ALTO", "REMISSÃO/FAV-91", onde,
                f"Rótulo cru no texto: {sorted(crus)[:4]}.",
                "Resolver a referência para 'Tabela N', 'Figura N' ou o nome da "
                "seção antes de entregar.")

        # título de seção solto dentro de parênteses, depois de vírgula
        for m in re.finditer(r"\(([^()]{0,60}),\s*([A-ZÁÉÍÓÚÂÊÔÃÕÇ][^()]{6,60}:"
                             r"[^()]{6,80})\)", txt):
            add("MÉDIO", "REMISSÃO/SWD-06", onde,
                f"Título de seção solto dentro de parênteses: "
                f"\u201c…{m.group(2)[:56]}…\u201d.",
                "Redigir a remissão ('ver a seção X') em vez de inserir o título "
                "nu, que o leitor lê como continuação da frase.")

        for m in re.finditer(r"\b(Tabela|Figura)\s+\1\s+\d", txt):
            add("ALTO", "REMISSÃO", onde,
                f"Palavra duplicada na chamada: “{m.group(0)[:40]}…”.",
                "A resolução da referência repetiu a palavra que já estava no "
                "texto; consumir a palavra anterior ao \\ref.")

        for m in re.finditer(r"\b(A|a)\s+(Figura|Tabela)\s+(?![\d~\\])", txt):
            trecho = txt[m.start():m.start() + 54].replace("\n", " ")
            add("MÉDIO", "REMISSÃO", onde,
                f"Chamada sem número: \u201c{trecho}…\u201d.",
                "A referência ficou órfã (o alvo saiu do documento) ou não foi "
                "resolvida.")


def check_log_linear() -> None:
    """Coeficiente em log ao lado do percentual: a diferença precisa de nota.

    Com a variável dependente em logaritmo, a variação percentual é
    (e^b − 1)×100, não b×100. Quem confere pela conta linear encontra outro
    número e supõe erro — aconteceu na revisão. A checagem confirma que a
    conversão exponencial explica o par e exige a explicação no documento.
    """
    import math

    alvo = ROOT / "tcc_normas.tex"
    if not alvo.exists():
        return
    txt = read(alvo)

    # pares "beta = -0,2123" ... "19,1%" na mesma vizinhança
    pares = []
    for m in re.finditer(r"hat\\beta[^$]{0,60}=\s*(−|-)?\s*0\{?,\}?(\d{3,4})", txt):
        b = -float("0." + m.group(2))
        viz = txt[max(0, m.start() - 320):m.start() + 320]
        for pm in re.finditer(r"(\d{1,2})\{?,\}?(\d)\\%", viz):
            pctv = float(f"{pm.group(1)}.{pm.group(2)}")
            exponencial = abs((math.exp(b) - 1) * 100)
            linear = abs(b) * 100
            if abs(pctv - exponencial) < 0.15 and abs(linear - exponencial) > 0.5:
                pares.append((b, pctv, linear))
                break

    if not pares:
        return
    explica = re.search(r"e\^\{?\\?hat?\\?beta\}?\s*-\s*1|\(e\^|exponencial",
                        txt) or "log-pontos" in txt and "\\footnote" in txt
    if not explica:
        b, pctv, linear = pares[0]
        add("MÉDIO", "LEITURA-LOG/SWD-06", "tcc_normas.tex",
            f"O texto traz \u03b2 = {b:.4f} e {pctv}%, que só fecham pela conversão "
            f"exponencial; pela leitura linear seriam {linear:.1f}%. "
            f"{len(pares)} par(es) nessa situação.",
            "Explicar uma vez, em nota de rodapé, que a variação percentual é "
            "(e^\u03b2 \u2212 1)\u00d7100 — sem isso o leitor confere pela conta linear e "
            "supõe erro.")


# ──────────────────────────────────────────────────────────────────────────────
def main() -> int:
    # stdout do Windows costuma ser cp1252 — forçar UTF-8 para não quebrar no hook
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--quiet", action="store_true")
    ap.add_argument("--no-save", action="store_true")
    ap.add_argument("--arquivo", default="")
    args = ap.parse_args()

    tex = read_tex()
    linhas = read(TEX).splitlines()          # .tex bruto: numeração real das linhas
    if tex:
        check_ml(tex); check_glmm(tex); check_qr(tex); check_hlm(tex); check_ob(tex)
        check_gap_interno(tex); check_escopo(linhas); check_causal(linhas); check_visual(tex)
    else:
        add("ALTO", "ESCOPO", str(TEX), "relatorio_tcc_enxuto.tex não encontrado.", "")
    check_scripts(); check_slides(); check_entregaveis()
    check_remissoes(); check_log_linear()
    if tex:
        check_carga_cognitiva(tex)

    ordem = {"ALTO": 0, "MÉDIO": 1, "BAIXO": 2, "INFO": 3}
    achados.sort(key=lambda a: ordem[a["nivel"]])
    n = {k: sum(1 for a in achados if a["nivel"] == k) for k in ordem}

    hoje = dt.datetime.now()
    md = [f"# Verificação mecânica — {hoje:%Y-%m-%d %H:%M}", "",
          f"Gatilho: `{args.arquivo or 'execução manual'}`  ",
          f"Achados: **{n['ALTO']} ALTO**, {n['MÉDIO']} MÉDIO, {n['BAIXO']} BAIXO, {n['INFO']} INFO.", "",
          "| Nível | Critério | Onde | Achado | Correção |", "|---|---|---|---|---|"]
    for a in achados:
        md.append(f"| {a['nivel']} | {a['crit']} | {a['onde']} | {a['msg']} | {a['fix']} |")
    md += ["", "_Checagens mecânicas apenas. Rode `/revisao-livros` para a revisão qualitativa completa "
           "(MHE, Fávero, Knaflic, Alencar)._"]
    if not args.no_save:
        out = ROOT / "tcc" / "revisoes" / f"verificacao_{hoje:%Y-%m-%d}.md"
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text("\n".join(md), encoding="utf-8")
    if args.quiet:
        for a in achados:
            if a["nivel"] in ("ALTO", "MÉDIO"):
                print(f"[{a['nivel']}] {a['crit']} — {a['onde']}: {a['msg']}")
    else:
        print("\n".join(md))
    print(f"\nRESUMO: {n['ALTO']} ALTO | {n['MÉDIO']} MÉDIO | {n['BAIXO']} BAIXO | {n['INFO']} INFO")
    return 0


if __name__ == "__main__":
    sys.exit(main())
